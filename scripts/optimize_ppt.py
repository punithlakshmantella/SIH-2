import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

def backup_and_load(ppt_path):
    backup_path = ppt_path.replace(".pptx", "_Original_Backup.pptx")
    if not os.path.exists(backup_path):
        shutil.copy2(ppt_path, backup_path)
        print(f"[+] Backup created at: {backup_path}")
    return Presentation(ppt_path)

def inspect_slide_shapes(prs):
    for i, slide in enumerate(prs.slides, 1):
        print(f"\n--- SLIDE {i} ---")
        for j, shape in enumerate(slide.shapes):
            if shape.has_text_frame:
                txt = " ".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
                print(f"Shape {j} ({shape.name}): {txt[:80]}...")

if __name__ == "__main__":
    ppt_path = r"C:\Users\Punit\OneDrive\Pictures\Route Riders.pptx"
    prs = backup_and_load(ppt_path)
    inspect_slide_shapes(prs)
